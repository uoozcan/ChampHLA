#!/usr/bin/env python3
"""Build normalised CIWD 3.0.0 lookup tables from the published IHIW xlsx supplements.

Source (downloaded into assets/ciwd_raw/ — see README_CIWD.md for URLs):
  * CIWD_Palleles-all_loci-2020320-ihws-website.xlsx   -> P group, two-field (primary)
  * SupTbl5Rev-Ggroup-CIWD-20200323_ihws-website.xlsx  -> G group (fallback)

Emits:
  * assets/ciwd_3.0.0.tsv          (allele, designation, ciwd_highest, ciwd_total,
                                     ciwd_AFA..ciwd_UNK, cwd2_category)
  * assets/ciwd_3.0.0_ggroup.tsv   (same columns, keyed on the G-group allele)

Category codes in the source are C / I / WD / (blank). We keep them verbatim here;
bin/ciwd.py maps them to common / intermediate / well_documented / not_ciwd, and treats
alleles absent from the table as `unknown`. Reproducible: re-run after re-downloading.
"""
import re
import sys
from pathlib import Path

import openpyxl

ASSETS = Path(__file__).resolve().parent
RAW = ASSETS / "ciwd_raw"

GROUPS = ["AFA", "API", "EURO", "MENA", "HIS", "NAM", "UNK"]
ALLELE_RE = re.compile(r"^[A-Z0-9]+\*\d")


def clean_allele(value):
    """`A*01:04N c` -> `A*01:04N`; `A*01:01:01G d` -> `A*01:01:01G`."""
    if value is None:
        return ""
    text = str(value).strip()
    # drop a trailing footnote marker: whitespace + one-or-more lowercase letters
    text = re.sub(r"\s+[a-z]+$", "", text)
    return text.replace(" ", "")


def find_header(ws_rows):
    """Return (header_index, {colname: idx}) for the row carrying 'Total' and 'AFA'."""
    for idx, row in enumerate(ws_rows):
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "Total" in cells and "AFA" in cells:
            mapping = {name: i for i, name in enumerate(cells)}
            return idx, mapping
    raise ValueError("no header row with Total/AFA found")


def cell(row, idx):
    if idx is None or idx >= len(row) or row[idx] is None:
        return ""
    return str(row[idx]).strip()


def parse_workbook(path, is_ggroup):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in wb.sheetnames:
        if sheet.lower().startswith("reference"):
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        try:
            hidx, cols = find_header(rows)
        except ValueError:
            continue
        allele_col = 0
        desig_col = cols.get("Nomenclature Designation c", cols.get("Nomenclature Designation a", 1))
        highest_col = cols.get("Highest Frequency")
        total_col = cols.get("Total")
        group_cols = {g: cols.get(g) for g in GROUPS}
        # CWD 2.0 category column exists only in the P-group table (label "Category" @ 2.0.0 CWD)
        cwd2_col = None
        if not is_ggroup:
            # three "Category" columns sit in the header row (2.0.0 CWD, EFI, China);
            # the first one after Total is CWD 2.0.0 (see the "2.0.0 CWD" super-header)
            cat_cols = [i for i, name in enumerate(rows[hidx])
                        if str(name).strip() == "Category" and i > (total_col or 0)]
            cwd2_col = cat_cols[0] if cat_cols else None
        for row in rows[hidx + 1:]:
            allele = clean_allele(cell(row, allele_col))
            if not ALLELE_RE.match(allele):
                continue
            records.append({
                "allele": allele,
                "designation": cell(row, desig_col),
                "ciwd_highest": cell(row, highest_col),
                "ciwd_total": cell(row, total_col),
                **{"ciwd_%s" % g: cell(row, group_cols[g]) for g in GROUPS},
                "cwd2_category": cell(row, cwd2_col) if cwd2_col is not None else "",
            })
    return records


def write_tsv(records, path):
    header = ["allele", "designation", "ciwd_highest", "ciwd_total"] \
        + ["ciwd_%s" % g for g in GROUPS] + ["cwd2_category"]
    seen = {}
    for rec in records:
        # prefer an expressed (P/G) row over a null row on key collision
        key = rec["allele"]
        if key not in seen or (rec["ciwd_total"] and not seen[key]["ciwd_total"]):
            seen[key] = rec
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\t".join(header) + "\n")
        for allele in sorted(seen):
            rec = seen[allele]
            fh.write("\t".join(rec.get(col, "") for col in header) + "\n")
    return len(seen)


def main():
    p_path = RAW / "CIWD_Palleles-all_loci-2020320-ihws-website.xlsx"
    g_path = RAW / "SupTbl5Rev-Ggroup-CIWD-20200323_ihws-website.xlsx"
    if not p_path.exists():
        sys.exit("Missing %s — download the CIWD 3.0.0 xlsx first (see README_CIWD.md)." % p_path)
    n_p = write_tsv(parse_workbook(p_path, is_ggroup=False), ASSETS / "ciwd_3.0.0.tsv")
    print("wrote assets/ciwd_3.0.0.tsv  (%d alleles)" % n_p)
    if g_path.exists():
        n_g = write_tsv(parse_workbook(g_path, is_ggroup=True), ASSETS / "ciwd_3.0.0_ggroup.tsv")
        print("wrote assets/ciwd_3.0.0_ggroup.tsv  (%d alleles)" % n_g)


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Build a Word .docx from the ChampHLA manuscript markdown, embedding figures.

The manuscript markdown (docs/CHAMPHLA_MANUSCRIPT_V3.md) is single-line-per-paragraph
and uses a small, regular subset of Markdown:
  - `# ` document title, `## ` section (Heading 1), `### ` subsection (Heading 2)
  - pipe tables with a `|---|` separator row
  - inline **bold**, *italic*, `code`
  - figure caption paragraphs of the form
        **Figure N.** *title* ... *(Source: `figures_final/<stem>`)*
    -> the caption is rendered (Source tail stripped) and the PNG embedded below it,
       resolved as <analysis-dir>/<source-path>.png (the source path already names its
       subdir, e.g. figures_final/ or figures_final_candidate/).
  - `---` horizontal rules (skipped), numbered `N. ` lines (rendered as paragraphs).

Supplementary figures (S1-S7) are referenced only in prose (no Source caption) and are
therefore not embedded — this build is faithful to the markdown's Source refs.

Usage:
  python3 bin/build_manuscript_docx.py <input.md> <output.docx> [--analysis-dir DIR]
"""
import argparse
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SOURCE_RE = re.compile(r"\(Source:\s*`([^`]+)`")
# bold first so ** wins over * at the same position; then italic, then code span
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*].*?\*|`[^`]+`)")
BASE_FONT = "Arial"
FIG_WIDTH_IN = 6.0


def add_inline_runs(paragraph, text):
    """Add `text` to `paragraph`, honouring **bold**, *italic*, `code`."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("**") and tok.endswith("**"):
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            run = paragraph.add_run(tok[1:-1])
            run.font.name = "Consolas"
        else:  # *italic*
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = BASE_FONT
    style.font.size = Pt(11)


def add_table(doc, rows):
    """rows = list of markdown table lines (incl. the |---| separator)."""
    parsed = []
    for ln in rows:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        parsed.append(cells)
    # drop the separator row (second line, all dashes)
    body = [parsed[0]] + parsed[2:] if len(parsed) >= 2 else parsed
    ncol = max(len(r) for r in body)
    table = doc.add_table(rows=len(body), cols=ncol)
    table.style = "Light Grid Accent 1"
    for i, row in enumerate(body):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ""
            para = cell.paragraphs[0]
            add_inline_runs(para, row[j] if j < len(row) else "")
            if i == 0:  # header bold
                for run in para.runs:
                    run.bold = True
    doc.add_paragraph()


def embed_figure(doc, analysis_dir, source_path, warnings):
    png = os.path.join(analysis_dir, source_path + ".png")
    if not os.path.isfile(png):
        warnings.append("MISSING figure PNG: " + png)
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(png, width=Inches(FIG_WIDTH_IN))


def build(md_path, docx_path, analysis_dir):
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    set_base_style(doc)
    warnings = []

    i = 0
    n = len(lines)
    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped == "---":  # horizontal rule
            i += 1
            continue

        # table block: consecutive lines starting with '|'
        if stripped.startswith("|"):
            block = []
            while i < n and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue

        # headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue

        # figure caption paragraph with a (Source: `...`) ref
        m = SOURCE_RE.search(stripped)
        if m and (stripped.startswith("**Figure") or stripped.startswith("**Supplementary Figure")):
            source_path = m.group(1)
            # strip the trailing *(Source: ...)* provenance note from the caption
            caption = re.sub(r"\s*\*?\(Source:.*?\)\*?\s*$", "", stripped)
            para = doc.add_paragraph()
            add_inline_runs(para, caption)
            embed_figure(doc, analysis_dir, source_path, warnings)
            i += 1
            continue

        # ordinary paragraph (incl. numbered 'N. ' lines)
        para = doc.add_paragraph()
        add_inline_runs(para, stripped)
        i += 1

    doc.save(docx_path)
    return warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("input_md")
    ap.add_argument("output_docx")
    ap.add_argument("--analysis-dir", default=None,
                    help="root for resolving figure Source paths "
                         "(default: <repo>/analysis, inferred from this script)")
    args = ap.parse_args()

    analysis_dir = args.analysis_dir
    if analysis_dir is None:
        repo = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        analysis_dir = os.path.join(repo, "analysis")

    warnings = build(args.input_md, args.output_docx, analysis_dir)
    print("[docx] wrote", args.output_docx)
    if warnings:
        print("[docx] {} warning(s):".format(len(warnings)))
        for w in warnings:
            print("  -", w)
    else:
        print("[docx] all referenced figures embedded.")


if __name__ == "__main__":
    main()

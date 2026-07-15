#!/usr/bin/env python3.11
"""
Convert CHAMPHLA_MANUSCRIPT_V3.md to a Word document with embedded figures.

Usage:
    python3.11 bin/manuscript_to_docx.py [--input docs/CHAMPHLA_MANUSCRIPT_V3.md] [--output docs/CHAMPHLA_MANUSCRIPT_V3.docx]
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


REPO_ROOT = Path(__file__).resolve().parent.parent
FIGURES_DIR = REPO_ROOT / "analysis" / "figures_final"

FIGURE_MAP = {
    "figure_1_workflow_architecture": "figure_1_workflow_architecture.png",
    "figure_2_accuracy_comparison": "figure_2_accuracy_comparison.png",
    "figure_3_per_gene_gains": "figure_3_per_gene_gains.png",
    "figure_4_confidence_calibration": "figure_4_confidence_calibration.png",
    "figure_5_abstention_tradeoff": "figure_5_abstention_tradeoff.png",
    "figure_6_discordance_taxonomy": "figure_6_discordance_taxonomy.png",
    "figure_7_confidence_weights": "figure_7_confidence_weights.png",
    "figure_8_computational_performance": "figure_8_computational_performance.png",
    "figure_09_bimodal_per_gene": "figure_09_bimodal_per_gene.png",
    "figure_10_trimodal_comparison": "figure_10_trimodal_comparison.png",
    "figure_s1_per_gene_accuracy": "figure_s1_per_gene_accuracy.png",
    "figure_s2_calibration_heatmap": "figure_s2_calibration_heatmap.png",
    "figure_s3_resolution_comparison": "figure_s3_resolution_comparison.png",
}


def setup_styles(doc):
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    pf = style.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    for level in range(1, 5):
        sname = f"Heading {level}"
        if sname in doc.styles:
            hs = doc.styles[sname]
            hs.font.name = "Arial"
            hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
            sizes = {1: 16, 2: 14, 3: 12, 4: 11}
            hs.font.size = Pt(sizes.get(level, 11))
            hs.font.bold = True


def add_inline_formatting(paragraph, text):
    bold_italic_re = re.compile(r"\*\*\*(.+?)\*\*\*")
    bold_re = re.compile(r"\*\*(.+?)\*\*")
    italic_re = re.compile(r"\*(.+?)\*")
    code_re = re.compile(r"`(.+?)`")

    combined = re.compile(
        r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|\*.+?\*|`.+?`)"
    )
    parts = combined.split(text)

    for part in parts:
        if not part:
            continue
        if bold_italic_re.fullmatch(part):
            run = paragraph.add_run(bold_italic_re.fullmatch(part).group(1))
            run.bold = True
            run.italic = True
        elif bold_re.fullmatch(part):
            run = paragraph.add_run(bold_re.fullmatch(part).group(1))
            run.bold = True
        elif italic_re.fullmatch(part):
            run = paragraph.add_run(italic_re.fullmatch(part).group(1))
            run.italic = True
        elif code_re.fullmatch(part):
            run = paragraph.add_run(code_re.fullmatch(part).group(1))
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def try_embed_figure(doc, line):
    source_match = re.search(r"\(Source:\s*`figures_final/(.+?)`\)", line)
    if not source_match:
        return False
    stem = source_match.group(1)
    png_name = FIGURE_MAP.get(stem)
    if not png_name:
        for k, v in FIGURE_MAP.items():
            if stem in k or k in stem:
                png_name = v
                break
    if not png_name:
        return False
    fig_path = FIGURES_DIR / png_name
    if not fig_path.exists():
        return False

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(fig_path), width=Inches(6.0))

    caption_text = line.strip()
    caption_text = re.sub(r"\s*\*\(Source:.*?\)\*\s*$", "", caption_text)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption_p.paragraph_format.space_before = Pt(4)
    caption_p.paragraph_format.space_after = Pt(12)
    add_inline_formatting(caption_p, caption_text)
    for run in caption_p.runs:
        run.font.size = Pt(9)
    return True


def add_table(doc, header_line, rows):
    headers = [c.strip() for c in header_line.split("|") if c.strip()]
    n_cols = len(headers)

    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline_formatting(p, h)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = "Arial"

    for row_text in rows:
        cells = [c.strip() for c in row_text.split("|") if c.strip()]
        if len(cells) != n_cols:
            continue
        row = table.add_row()
        for i, c in enumerate(cells):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_formatting(p, c)
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.name = "Arial"

    doc.add_paragraph()


def convert(md_path, docx_path):
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    i = 0
    in_code_block = False
    code_lines = []
    table_header = None
    table_rows = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code_block:
                code_text = "\n".join(code_lines)
                p = doc.add_paragraph()
                run = p.add_run(code_text)
                run.font.name = "Courier New"
                run.font.size = Pt(8.5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                in_code_block = False
                code_lines = []
            else:
                if table_header:
                    add_table(doc, table_header, table_rows)
                    table_header = None
                    table_rows = []
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("|") and "|" in line[1:]:
            cells = [c.strip() for c in line.split("|") if c.strip()]
            if all(set(c) <= set("-: ") for c in cells):
                i += 1
                continue
            if table_header is None:
                table_header = line
            else:
                table_rows.append(line)
            i += 1
            continue
        else:
            if table_header:
                add_table(doc, table_header, table_rows)
                table_header = None
                table_rows = []

        if line.strip() == "---":
            doc.add_paragraph()
            i += 1
            continue

        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=4)
            i += 1
            continue

        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if "*(Source: `figures_final/" in stripped:
            if try_embed_figure(doc, stripped):
                i += 1
                continue

        p = doc.add_paragraph()
        add_inline_formatting(p, stripped)
        i += 1

    if table_header:
        add_table(doc, table_header, table_rows)

    doc.save(str(docx_path))
    print(f"Saved: {docx_path} ({Path(docx_path).stat().st_size / 1024:.0f} KB)")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input",
        default=str(REPO_ROOT / "docs" / "CHAMPHLA_MANUSCRIPT_V3.md"),
    )
    parser.add_argument(
        "--output",
        default=str(REPO_ROOT / "docs" / "CHAMPHLA_MANUSCRIPT_V3.docx"),
    )
    args = parser.parse_args()
    convert(args.input, args.output)


if __name__ == "__main__":
    main()
